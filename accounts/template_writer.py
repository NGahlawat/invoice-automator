from openpyxl import load_workbook


def update_nourish_details(
    file_path,
    excel_row,
    nourish_id=None,
    nourish_hours=None,
    visit_breakdown=None,
    hours_difference=None,
    validation_status=None,
    validation_notes=None
):

    workbook = load_workbook(
        file_path
    )

    sheet = workbook[
        "Monthly Accounts"
    ]

    # Nourish Client ID
    if nourish_id is not None:

        sheet.cell(
            row=excel_row,
            column=13
        ).value = str(
            nourish_id
        )

    # Nourish Billable Hours
    if nourish_hours is not None:

        sheet.cell(
            row=excel_row,
            column=14
        ).value = float(
            nourish_hours
        )

    # Nourish Visit Breakdown
    if visit_breakdown is not None:

        sheet.cell(
            row=excel_row,
            column=15
        ).value = visit_breakdown

    # Hours Difference
    if hours_difference is not None:

        sheet.cell(
            row=excel_row,
            column=16
        ).value = float(
            hours_difference
        )

    # Validation Status
    if validation_status is not None:

        sheet.cell(
            row=excel_row,
            column=17
        ).value = validation_status

    # Validation Notes
    if validation_notes is not None:

        sheet.cell(
            row=excel_row,
            column=18
        ).value = validation_notes

    workbook.save(
        file_path
    )