from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side
)
from config import OUTPUT_FOLDER


LOGO_FILE = "assets/logo.png"


def create_invoice(
    user,
    invoice_lines,
    total,
    invoice_number,
    invoice_date,
    period,
    company_settings
):

    workbook = Workbook()

    ws = workbook.active
    ws.title = "Invoice"

    ws.sheet_view.showGridLines = False


    # --------------------------------------------------
    # COLUMN WIDTHS
    # --------------------------------------------------

    ws.column_dimensions["A"].width = 27
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18


    # --------------------------------------------------
    # ROW HEIGHTS
    # --------------------------------------------------

    for row in range(1, 50):
        ws.row_dimensions[row].height = 16

    ws.row_dimensions[1].height = 35


    # --------------------------------------------------
    # FONTS
    # --------------------------------------------------

    title_font = Font(
        name="Arial",
        size=18,
        bold=True
    )

    normal_font = Font(
        name="Arial",
        size=10
    )

    bold_font = Font(
        name="Arial",
        size=10,
        bold=True
    )

    small_font = Font(
        name="Arial",
        size=8
    )

    terms_font = Font(
        name="Arial",
        size=8,
        underline="single"
    )

    green_font = Font(
        name="Arial",
        size=11,
        bold=True,
        color="00A000"
    )

    white_bold_font = Font(
        name="Arial",
        size=10,
        bold=True,
        color="FFFFFF"
    )


    # --------------------------------------------------
    # COLOURS / BORDERS
    # --------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="808080"
    )

    thin_black = Side(
        style="thin",
        color="000000"
    )

    top_border = Border(
        top=thin_black
    )


    # --------------------------------------------------
    # TITLE
    # --------------------------------------------------

    ws["A1"] = "INVOICE"
    ws["A1"].font = title_font


    # --------------------------------------------------
    # LOGO
    # --------------------------------------------------

    if Path(LOGO_FILE).exists():

        logo = Image(LOGO_FILE)

        logo.width = 230
        logo.height = 65

        ws.add_image(
            logo,
            "C1"
        )


    # --------------------------------------------------
    # COMPANY / CLIENT INFORMATION
    # --------------------------------------------------

    ws.merge_cells(
        "A5:B5"
    )

    ws["A5"] = (
        company_settings.get(
            "company_name"
        )
        or ""
    )

    ws["A5"].font = bold_font


    ws.merge_cells(
        "A7:B7"
    )

    ws["A7"] = (
        "Client PIN - "
        + str(
            user.get(
                "reference"
            )
            or ""
        )
    )

    ws["A7"].font = bold_font


    # --------------------------------------------------
    # INVOICE INFORMATION
    # --------------------------------------------------

    ws["C5"] = "Invoice Date"
    ws["D5"] = invoice_date

    ws["C6"] = "Invoice No"
    ws["D6"] = invoice_number

    ws["C7"] = "PO Number"
    ws["D7"] = (
        user.get(
            "po_number"
        )
        or ""
    )


    for cell in [
        "C5",
        "C6",
        "C7",
        "D5",
        "D6",
        "D7"
    ]:

        ws[cell].font = normal_font


    # --------------------------------------------------
    # SUPPLIER ADDRESS
    # --------------------------------------------------

    ws["A10"] = "Supplier Address"
    ws["A10"].font = normal_font


    supplier_address = [
        company_settings.get(
            "company_name"
        ),
        company_settings.get(
            "supplier_address_1"
        ),
        company_settings.get(
            "supplier_address_2"
        ),
        company_settings.get(
            "supplier_postcode"
        )
    ]


    supplier_row = 11


    for line in supplier_address:

        if not line:
            continue

        ws.merge_cells(
            start_row=supplier_row,
            start_column=1,
            end_row=supplier_row,
            end_column=2
        )

        ws.cell(
            supplier_row,
            1
        ).value = line

        ws.cell(
            supplier_row,
            1
        ).font = normal_font

        supplier_row += 1


    # --------------------------------------------------
    # BILL TO
    # --------------------------------------------------

    ws["C9"] = "Bill To:-"
    ws["C9"].font = normal_font


    funding_type = str(
        user.get(
            "funding_type"
        )
        or ""
    ).lower()


    if funding_type == "private":

        bill_to = [
            user.get(
                "billing_address_1"
            ),
            user.get(
                "billing_address_2"
            ),
            user.get(
                "billing_city"
            ),
            user.get(
                "billing_postcode"
            )
        ]

    else:

        bill_to = [
            company_settings.get(
                "nhs_bill_to_1"
            ),
            company_settings.get(
                "nhs_bill_to_2"
            ),
            company_settings.get(
                "nhs_bill_to_3"
            ),
            company_settings.get(
                "nhs_bill_to_4"
            ),
            company_settings.get(
                "nhs_bill_to_5"
            ),
            company_settings.get(
                "nhs_bill_to_postcode"
            )
        ]


    bill_to = [
        line
        for line in bill_to
        if line
    ]


    bill_row = 9


    for line in bill_to:

        ws.merge_cells(
            start_row=bill_row,
            start_column=4,
            end_row=bill_row,
            end_column=5
        )

        ws.cell(
            bill_row,
            4
        ).value = line

        ws.cell(
            bill_row,
            4
        ).font = normal_font

        bill_row += 1


    # --------------------------------------------------
    # TABLE
    # --------------------------------------------------

    header_row = 18


    headings = [
        "Invoice Period",
        "DESCRIPTION",
        "Hours",
        "Rate",
        "Amount"
    ]


    for column, heading in enumerate(
        headings,
        start=1
    ):

        cell = ws.cell(
            header_row,
            column
        )

        cell.value = heading
        cell.font = white_bold_font
        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    current_row = (
        header_row
        + 1
    )

    previous_care_type = None


    for line in invoice_lines:

        care_type = line.get(
            "care_type"
        )


        if care_type != previous_care_type:

            ws.cell(
                current_row,
                2
            ).value = care_type

            ws.cell(
                current_row,
                2
            ).font = Font(
                name="Arial",
                size=10,
                bold=True,
                italic=True
            )

            ws.cell(
                current_row,
                2
            ).alignment = Alignment(
                horizontal="center"
            )

            current_row += 1

            previous_care_type = care_type


        ws.cell(
            current_row,
            1
        ).value = period


        ws.cell(
            current_row,
            2
        ).value = line.get(
            "description"
        )


        ws.cell(
            current_row,
            3
        ).value = float(
            line.get(
                "hours",
                0
            )
        )


        ws.cell(
            current_row,
            4
        ).value = float(
            line.get(
                "rate",
                0
            )
        )


        ws.cell(
            current_row,
            5
        ).value = float(
            line.get(
                "amount",
                0
            )
        )


        ws.cell(
            current_row,
            3
        ).number_format = "0.00"

        ws.cell(
            current_row,
            4
        ).number_format = "£0.00"

        ws.cell(
            current_row,
            5
        ).number_format = "£#,##0.00"


        ws.cell(
            current_row,
            3
        ).alignment = Alignment(
            horizontal="center"
        )

        ws.cell(
            current_row,
            4
        ).alignment = Alignment(
            horizontal="center"
        )

        ws.cell(
            current_row,
            5
        ).alignment = Alignment(
            horizontal="right"
        )


        current_row += 1


    # --------------------------------------------------
    # TOTAL
    # --------------------------------------------------

    total_row = max(
        current_row + 2,
        27
    )


    ws["D" + str(total_row)] = (
        "Total To Pay"
    )

    ws["E" + str(total_row)] = float(
        total
    )


    ws["D" + str(total_row)].font = bold_font
    ws["E" + str(total_row)].font = bold_font


    ws["D" + str(total_row)].border = top_border
    ws["E" + str(total_row)].border = top_border


    ws["D" + str(total_row)].alignment = Alignment(
        horizontal="center"
    )

    ws["E" + str(total_row)].alignment = Alignment(
        horizontal="right"
    )


    ws[
        "E" + str(total_row)
    ].number_format = "£#,##0.00"


    # --------------------------------------------------
    # PAYMENT SECTION
    # --------------------------------------------------

    footer_row = (
        total_row
        + 2
    )


    ws.merge_cells(
        start_row=footer_row,
        start_column=1,
        end_row=footer_row,
        end_column=2
    )

    ws.cell(
        footer_row,
        1
    ).value = (
        company_settings.get(
            "payment_heading"
        )
        or "Bank payment made to:"
    )


    # PAYMENT REFERENCE RIGHT SIDE

    ws.merge_cells(
        start_row=footer_row,
        start_column=3,
        end_row=footer_row + 1,
        end_column=5
    )

    ws.cell(
        footer_row,
        3
    ).value = (
        company_settings.get(
            "payment_reference_message"
        )
        or ""
    )

    ws.cell(
        footer_row,
        3
    ).font = bold_font

    ws.cell(
        footer_row,
        3
    ).alignment = Alignment(
        vertical="top",
        wrap_text=True
    )


    # COMPANY NAME

    ws.merge_cells(
        start_row=footer_row + 2,
        start_column=1,
        end_row=footer_row + 2,
        end_column=2
    )

    ws.cell(
        footer_row + 2,
        1
    ).value = (
        company_settings.get(
            "company_name"
        )
        or ""
    )

    ws.cell(
        footer_row + 2,
        1
    ).font = bold_font


    # BANK DETAILS

    account_number = (
        company_settings.get(
            "account_number"
        )
        or ""
    )

    sort_code = (
        company_settings.get(
            "sort_code"
        )
        or ""
    )


    ws.merge_cells(
        start_row=footer_row + 4,
        start_column=1,
        end_row=footer_row + 4,
        end_column=3
    )

    ws.cell(
        footer_row + 4,
        1
    ).value = (
        f"Account No- {account_number}, "
        f"Sort code- {sort_code}"
    )

    ws.cell(
        footer_row + 4,
        1
    ).font = bold_font


    # --------------------------------------------------
    # CONTACT MESSAGE
    # --------------------------------------------------

    contact_row = (
        footer_row
        + 6
    )


    ws.merge_cells(
        start_row=contact_row,
        start_column=1,
        end_row=contact_row,
        end_column=5
    )


    ws.cell(
        contact_row,
        1
    ).value = (
        company_settings.get(
            "invoice_contact_message"
        )
        or ""
    )


    ws.cell(
        contact_row,
        1
    ).font = small_font


    ws.cell(
        contact_row,
        1
    ).alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


    # --------------------------------------------------
    # TERMS
    # --------------------------------------------------

    terms_row = (
        contact_row
        + 1
    )


    ws.merge_cells(
        start_row=terms_row,
        start_column=1,
        end_row=terms_row + 1,
        end_column=5
    )


    payment_terms = (
        company_settings.get(
            "payment_terms"
        )
        or ""
    )


    ws.cell(
        terms_row,
        1
    ).value = (
        "TERMS: "
        + payment_terms
    )


    ws.cell(
        terms_row,
        1
    ).font = terms_font


    ws.cell(
        terms_row,
        1
    ).alignment = Alignment(
        horizontal="center",
        vertical="top",
        wrap_text=True
    )


    # --------------------------------------------------
    # APPRECIATION
    # --------------------------------------------------

    appreciation_row = (
        terms_row
        + 3
    )


    ws.merge_cells(
        start_row=appreciation_row,
        start_column=1,
        end_row=appreciation_row,
        end_column=5
    )


    ws.cell(
        appreciation_row,
        1
    ).value = (
        company_settings.get(
            "appreciation_message"
        )
        or ""
    )


    ws.cell(
        appreciation_row,
        1
    ).font = green_font


    ws.cell(
        appreciation_row,
        1
    ).alignment = Alignment(
        horizontal="center"
    )


    # --------------------------------------------------
    # PRINT SETTINGS
    # --------------------------------------------------

    ws.print_area = (
        f"A1:E{appreciation_row}"
    )


    ws.page_setup.orientation = "portrait"

    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    ws.sheet_properties.pageSetUpPr.fitToPage = True


    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.3
    ws.page_margins.bottom = 0.3


    # --------------------------------------------------
    # SAVE
    # --------------------------------------------------

    output_file = (
        OUTPUT_FOLDER
        / f"{invoice_number}.xlsx"
    )


    workbook.save(
        output_file
    )


    return output_file